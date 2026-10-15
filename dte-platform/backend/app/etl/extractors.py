"""Extract tabular data (and free text) from DOCX, PDF, PPTX, images, Excel, CSV, JSON and XML."""
from __future__ import annotations

import pathlib
import re
import shutil

import pandas as pd

SUPPORTED = {".csv": "csv", ".xlsx": "excel", ".xls": "excel", ".xlsm": "excel", ".json": "json", ".xml": "xml",
             ".docx": "docx", ".pdf": "pdf", ".pptx": "pptx", ".png": "image", ".jpg": "image", ".jpeg": "image",
             ".tif": "image", ".tiff": "image", ".bmp": "image"}


class Extraction:
    def __init__(self, file_type: str):
        self.file_type = file_type
        self.tables: list[dict] = []   # {"name": str, "df": DataFrame}
        self.text: str = ""
        self.log: list[str] = []

    def add(self, name: str, df: pd.DataFrame):
        df = repair_header(df, self.log, name)
        df = df.dropna(how="all").dropna(axis=1, how="all")
        if not df.empty:
            self.tables.append({"name": name, "df": df.reset_index(drop=True)})


def detect_type(filename: str) -> str:
    ext = pathlib.Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type {ext}. Supported: {', '.join(sorted(SUPPORTED))}")
    return SUPPORTED[ext]


def _clean_col(c, i) -> str:
    c = re.sub(r"\s+", " ", str(c)).strip()
    if not c or c.lower().startswith("unnamed") or c.lower() == "nan":
        return f"column_{i+1}"
    return c


def repair_header(df: pd.DataFrame, log: list, name: str) -> pd.DataFrame:
    """If the real header is not the first row (title rows above it, common in Excel exports), find it."""
    if df.empty:
        return df
    generic = sum(1 for c in df.columns if str(c).lower().startswith("unnamed") or isinstance(c, int))
    if generic >= max(1, len(df.columns) // 2):
        for i in range(min(10, len(df))):
            row = df.iloc[i]
            vals = [v for v in row if pd.notna(v) and str(v).strip()]
            if len(vals) >= max(2, int(len(df.columns) * 0.6)) and all(isinstance(v, str) for v in vals):
                df = df.iloc[i + 1:].copy()
                df.columns = [str(v) if pd.notna(v) else None for v in row]
                log.append(f"{name}: header found on row {i+1}, {i} title row(s) dropped")
                break
    cols, seen = [], {}
    for i, c in enumerate(df.columns):
        c2 = _clean_col(c, i)
        if c2 in seen:
            seen[c2] += 1
            c2 = f"{c2}_{seen[c2]}"
        else:
            seen[c2] = 0
        cols.append(c2)
    df.columns = cols
    return df


def _rows_to_df(rows: list[list]) -> pd.DataFrame | None:
    rows = [[(c or "").strip() if isinstance(c, str) or c is None else c for c in r] for r in rows if r and any(r)]
    if len(rows) < 2:
        return None
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return pd.DataFrame(rows[1:], columns=rows[0]).replace("", None)


def extract(path: str | pathlib.Path, filename: str | None = None) -> Extraction:
    path = pathlib.Path(path)
    ftype = detect_type(filename or path.name)
    ex = Extraction(ftype)
    if ftype == "csv":
        ex.add("csv", pd.read_csv(path, sep=None, engine="python"))
    elif ftype == "excel":
        for sheet, df in pd.read_excel(path, sheet_name=None).items():
            ex.add(f"sheet:{sheet}", df)
    elif ftype == "json":
        import json
        data = json.loads(path.read_text())
        if isinstance(data, dict):
            key = next((k for k, v in data.items() if isinstance(v, list)), None)
            data = data[key] if key else [data]
        ex.add("json", pd.json_normalize(data))
    elif ftype == "xml":
        try:
            df = pd.read_xml(path)
        except ImportError:
            df = pd.read_xml(path, parser="etree")
        ex.add("xml", df)
    elif ftype == "docx":
        import docx
        d = docx.Document(str(path))
        for i, t in enumerate(d.tables):
            df = _rows_to_df([[c.text for c in r.cells] for r in t.rows])
            if df is not None:
                ex.add(f"table:{i+1}", df)
        ex.text = "\n".join(p.text for p in d.paragraphs if p.text.strip())
    elif ftype == "pptx":
        from pptx import Presentation
        prs = Presentation(str(path))
        texts = []
        for si, slide in enumerate(prs.slides, 1):
            for sh in slide.shapes:
                if sh.has_table:
                    df = _rows_to_df([[c.text for c in r.cells] for r in sh.table.rows])
                    if df is not None:
                        ex.add(f"slide{si}:table", df)
                elif sh.has_text_frame and sh.text_frame.text.strip():
                    texts.append(sh.text_frame.text)
        ex.text = "\n".join(texts)
    elif ftype == "pdf":
        import pdfplumber
        texts = []
        with pdfplumber.open(str(path)) as pdf:
            for pi, page in enumerate(pdf.pages, 1):
                found = page.extract_tables()
                if not found:  # borderless tables: fall back to text-alignment detection
                    found = [t for t in page.extract_tables({"vertical_strategy": "text", "horizontal_strategy": "text"})
                             if t and len(t) >= 2 and len(t[0]) >= 2]
                for ti, tbl in enumerate(found, 1):
                    df = _rows_to_df(tbl)
                    if df is not None:
                        ex.add(f"page{pi}:table{ti}", df)
                texts.append(page.extract_text() or "")
        ex.text = "\n".join(texts)
        if not ex.tables and not ex.text.strip():
            ex.log.append("PDF has no text layer — run OCR by uploading page images")
    elif ftype == "image":
        if not shutil.which("tesseract"):
            raise RuntimeError("Image OCR needs the tesseract binary (apt install tesseract-ocr).")
        import pytesseract
        from PIL import Image
        img = Image.open(path)
        ex.text = pytesseract.image_to_string(img)
        df = _ocr_table_from_boxes(pytesseract.image_to_data(img, output_type=pytesseract.Output.DATAFRAME))
        if df is None:
            df = _ocr_table(ex.text)
        if df is not None:
            ex.add("ocr:table", df)
        ex.log.append("OCR applied (Tesseract)")
    # merge tables with identical columns (e.g. multi-page PDF table)
    ex.tables = _merge_same_shape(ex.tables, ex.log)
    return ex


def _ocr_table(text: str) -> pd.DataFrame | None:
    lines = [re.split(r"\s{2,}|\t|\|", ln.strip()) for ln in text.splitlines() if ln.strip()]
    lines = [[c for c in ln if c] for ln in lines]
    if not lines:
        return None
    width = max(set(len(l) for l in lines), key=lambda w: sum(1 for l in lines if len(l) == w))
    rows = [l for l in lines if len(l) == width]
    return _rows_to_df(rows) if width >= 2 and len(rows) >= 3 else None


def _ocr_table_from_boxes(data: pd.DataFrame):
    """Rebuild table cells from OCR word boxes: words on a line separated by a wide gap start a new cell."""
    words = data[(data.conf.astype(float) > 30) & data.text.notna()]
    words = words[words.text.astype(str).str.strip() != ""]
    if words.empty:
        return None
    gap = max(12, float(words.height.median()) * 1.2)
    lines = []
    for _, ln in words.sort_values(["block_num", "par_num", "line_num", "left"]).groupby(["block_num", "par_num", "line_num"]):
        cells, cur, right = [], [], None
        for w in ln.itertuples():
            if right is not None and w.left - right > gap:
                cells.append(" ".join(cur))
                cur = []
            cur.append(str(w.text))
            right = w.left + w.width
        cells.append(" ".join(cur))
        lines.append((int(ln.top.min()), cells))
    lines.sort()
    rows = [c for _, c in lines]
    if not rows:
        return None
    width = max(set(len(r) for r in rows), key=lambda w: sum(1 for r in rows if len(r) == w))
    rows = [r for r in rows if len(r) == width]
    df = _rows_to_df(rows) if width >= 2 and len(rows) >= 3 else None
    return df


def _merge_same_shape(tables: list[dict], log: list) -> list[dict]:
    out: list[dict] = []
    for t in tables:
        match = next((o for o in out if list(o["df"].columns) == list(t["df"].columns)), None)
        if match:
            match["df"] = pd.concat([match["df"], t["df"]], ignore_index=True)
            log.append(f"{t['name']} appended to {match['name']} (same columns)")
        else:
            out.append(t)
    return out


def extract_bytes(data: bytes, filename: str, staging_dir: pathlib.Path) -> tuple[Extraction, pathlib.Path]:
    staging_dir.mkdir(parents=True, exist_ok=True)
    p = staging_dir / re.sub(r"[^A-Za-z0-9._-]", "_", filename)
    p.write_bytes(data)
    return extract(p, filename), p

