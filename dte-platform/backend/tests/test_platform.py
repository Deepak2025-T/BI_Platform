"""End-to-end tests against a real PostgreSQL (+pgvector) with a scripted fake LLM.

Needs: dte_app DB (pgvector) and finance_demo DB seeded by sample/seed_finance.py.
Run:   pytest -q tests/
"""
import io
import json
import os

import pandas as pd
import pytest
from fastapi.testclient import TestClient

from app.llm import embeddings
from app.llm.providers import LLM, set_llm_override
from app.main import app

FIN_DB = os.getenv("TEST_FINANCE_DB", "finance_demo")

SQL_BY_KEYWORD = [
    ("broken", "SELECT SUM(amount_typo) AS total FROM invoices"),
    ("revenue by month", "SELECT date_trunc('month', invoice_date)::date AS period, SUM(amount_untaxed) AS revenue "
                         "FROM invoices WHERE invoice_type='sales' GROUP BY 1 ORDER BY 1"),
    ("by region", "SELECT p.region, SUM(i.total_amount) AS sales FROM invoices i JOIN parties p ON p.party_id=i.party_id "
                  "WHERE i.invoice_type='sales' GROUP BY 1 ORDER BY 2 DESC"),
    ("by type and month", "SELECT date_trunc('month', invoice_date)::date AS period, invoice_type, SUM(amount_untaxed) AS amount "
                          "FROM invoices GROUP BY 1,2 ORDER BY 1"),
    ("total revenue", "SELECT SUM(amount_untaxed) AS total_revenue FROM invoices WHERE invoice_type='sales'"),
    ("delete", "DELETE FROM invoices"),
    ("top customers", "SELECT p.name AS customer, SUM(i.total_amount) AS sales FROM invoices i JOIN parties p USING (party_id) "
                      "WHERE invoice_type='sales' GROUP BY 1 ORDER BY 2 DESC LIMIT 10"),
]


class FakeLLM(LLM):
    name = "fake"

    def __init__(self):
        self.calls = []

    def complete(self, system, user, *, json_mode=False, tier="strong"):
        self.calls.append((system[:40], user[:80]))
        if system.startswith("You are the IntentAgent"):
            return json.dumps({"intent": "trend", "measures": ["revenue"], "dimensions": ["month"], "search_query": "invoices revenue"})
        if system.startswith("You are the ReportPlanner"):
            return json.dumps({"title": "Finance overview", "widgets": [
                {"slot": "kpi1", "title": "Revenue", "question": "total revenue"},
                {"slot": "trend", "title": "Revenue trend", "question": "revenue by month"},
                {"slot": "mix", "title": "Sales by region", "question": "sales by region"},
                {"slot": "detail", "title": "Top customers", "question": "top customers"},
            ]})
        if system.startswith("You are the LayoutAgent"):
            return json.dumps({"insights": ["Revenue grew steadily.", "Largest month was recent."]})
        if system.startswith("You are the SchemaAgent"):
            return json.dumps({"mapping": {
                "invoice_no": {"source": "Invoice #", "transform": "strip", "confidence": 0.95},
                "invoice_type": {"source": "Type", "transform": "lower", "confidence": 0.9},
                "party_id": {"source": "Customer", "transform": "none", "confidence": 0.9},
                "invoice_date": {"source": "Date", "transform": "to_date", "confidence": 0.9},
                "total_amount": {"source": "Amount (INR)", "transform": "to_number", "confidence": 0.9},
            }, "notes": "ok"})
        if system.startswith("You pick the best target table"):
            return json.dumps({"table": "invoices"})
        if system.startswith("You are the SQLAgent"):
            q = user.lower()
            if "the previous sql failed" in q:
                return json.dumps({"sql": "SELECT SUM(total_amount) AS total FROM invoices", "chart": {"type": "kpi"}})
            for k, sql in SQL_BY_KEYWORD:
                if k in q:
                    return json.dumps({"sql": sql, "explanation": k, "chart": {}})
            return json.dumps({"sql": "SELECT COUNT(*) AS n FROM invoices"})
        return "OK"


@pytest.fixture(scope="module")
def client():
    embeddings.use_fallback(True)
    fake = FakeLLM()
    set_llm_override(fake)
    with TestClient(app) as c:
        c.fake = fake
        yield c
    set_llm_override(None)


@pytest.fixture(scope="module")
def conn_id(client):
    for c in client.get("/api/connections").json():
        if c["name"] == "pytest-finance":
            client.delete(f"/api/connections/{c['id']}")
    r = client.post("/api/connections", json={"name": "pytest-finance", "dialect": "postgres", "host": "localhost",
                                               "port": 5432, "database": FIN_DB, "username": "postgres", "password": "postgres"})
    assert r.status_code == 200, r.text
    cid = r.json()["id"]
    assert client.post(f"/api/connections/{cid}/test").json()["ok"]
    r = client.post(f"/api/connections/{cid}/index", json={"glossary": {"revenue": "SUM(amount_untaxed) of sales invoices"}})
    assert r.json()["tables_indexed"] >= 8
    yield cid
    # keep the demo DB clean for the next run
    from sqlalchemy import create_engine, text
    eng = create_engine(f"postgresql+psycopg://postgres:postgres@localhost:5432/{FIN_DB}")
    with eng.begin() as c:
        c.execute(text("DELETE FROM invoices WHERE invoice_no LIKE 'ETL-T-%'"))
        c.execute(text("DELETE FROM accounts WHERE code = '9999'"))


def test_dialects(client):
    keys = {d["key"] for d in client.get("/api/dialects").json()}
    assert keys == {"postgres", "mysql", "mssql", "oracle", "db2", "sqlite"}


def test_trend_query(client, conn_id):
    r = client.post("/api/query", json={"connection_id": conn_id, "question": "Show revenue by month"}).json()
    assert r["error"] is None, r
    assert r["chart"]["type"] == "line" and r["chart"]["x"] == "period"
    assert len(r["rows"]) == 24
    assert r["kpis"][0]["change_pct"] is not None
    assert "invoices" in r["tables_used"]


def test_stacked_and_pie(client, conn_id):
    r = client.post("/api/query", json={"connection_id": conn_id, "question": "amount by type and month"}).json()
    assert r["chart"]["type"] == "stacked_column" and r["chart"]["series"] == "invoice_type"
    r = client.post("/api/query", json={"connection_id": conn_id, "question": "sales by region", "chart": "pie"}).json()
    assert r["chart"]["type"] == "pie" and r["chart"]["x"] == "region"


def test_autofix(client, conn_id):
    r = client.post("/api/query", json={"connection_id": conn_id, "question": "broken query please"}).json()
    assert r["attempts"] == 2 and r["error"] is None
    assert r["chart"]["type"] == "kpi"


def test_write_blocked(client, conn_id):
    r = client.post("/api/query/sql", json={"connection_id": conn_id, "sql": "DELETE FROM invoices"})
    assert r.status_code == 400 and "Only SELECT" in r.json()["detail"]
    r = client.post("/api/query/sql", json={"connection_id": conn_id, "sql": "SELECT 1; DROP TABLE invoices"})
    assert r.status_code == 400
    r = client.post("/api/query/sql", json={"connection_id": conn_id, "sql": "SELECT * FROM invoices", "max_rows": 7})
    assert len(r.json()["rows"]) == 7


def test_generate_save_publish(client, conn_id):
    assert len(client.get("/api/templates").json()) == 5
    spec = client.post("/api/reports/generate", json={"connection_id": conn_id, "prompt": "Finance overview FY", "template_id": "executive"}).json()
    assert len(spec["widgets"]) == 8
    trend = next(w for w in spec["widgets"] if w["id"] == "trend")
    assert trend["chart"]["type"] == "line" and trend["insights"]
    mix = next(w for w in spec["widgets"] if w["id"] == "mix")
    assert mix["chart"]["type"] == "donut"
    r = client.post("/api/reports", json={"title": spec["title"], "prompt": spec["prompt"], "connection_id": conn_id,
                                          "template_id": "executive", "spec": spec}).json()
    got = client.get(f"/api/reports/{r['id']}").json()
    assert got["spec"]["widgets"][1]["rows"]  # data re-queried on open
    assert client.post(f"/api/reports/{r['id']}/publish").json()["examples_saved"] >= 4


def test_forecast(client, conn_id):
    r = client.post("/api/forecast", json={"connection_id": conn_id, "metric": "revenue", "horizon": 6}).json()
    assert len(r["forecast"]) == 6 and r["model"].startswith("holt_winters")
    assert all(f["lower"] <= f["value"] <= f["upper"] for f in r["forecast"])


def test_etl_excel_flow(client, conn_id):
    df = pd.DataFrame({"Invoice #": [" ETL-T-001 ", "ETL-T-002"], "Type": ["SALES", "SALES"],
                       "Customer": ["GKW Consult", "acme industries"], "Date": ["2026-10-01", "05/10/2026"],
                       "Amount (INR)": ["₹1,18,000.00", "59000"]})
    buf = io.BytesIO()
    with pd.ExcelWriter(buf) as xw:
        pd.DataFrame([["Invoice register Oct 2026"], [None]]).to_excel(xw, index=False, header=False, startrow=0)
        df.to_excel(xw, index=False, startrow=2)
    r = client.post("/api/etl/extract", data={"connection_id": conn_id},
                    files={"file": ("register.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    run = r.json()
    assert run["tables"][0]["columns"][0] == "Invoice #"   # title rows repaired
    rid = run["id"]
    m = client.post(f"/api/etl/{rid}/map", json={"source_index": 0}).json()
    assert m["target_table"] == "invoices"
    t = client.post(f"/api/etl/{rid}/transform", json={}).json()
    assert t["status"] == "validated", t["issues"]
    rows = t["preview"]["rows"]
    assert rows[0]["invoice_no"] == "ETL-T-001" and rows[0]["total_amount"] == 118000.0 and rows[1]["party_id"]
    loaded = client.post(f"/api/etl/{rid}/load", json={"mode": "upsert"}).json()
    assert loaded["rows_loaded"] == 2
    # second time: duplicate warning, append blocked, upsert allowed
    t2 = client.post(f"/api/etl/{rid}/transform", json={}).json()
    assert any("already exist" in i["message"] for i in t2["issues"])
    assert client.post(f"/api/etl/{rid}/load", json={"mode": "append"}).status_code == 409
    assert client.post(f"/api/etl/{rid}/load", json={"mode": "upsert"}).status_code == 200


def test_etl_docx_pdf_xml(client, conn_id, tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Journal voucher")
    t = d.add_table(rows=3, cols=3)
    for i, row in enumerate([["Account", "Debit", "Credit"], ["4000", "0", "500"], ["1100", "500", "0"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    p = tmp_path / "jv.docx"
    d.save(p)
    r = client.post("/api/etl/extract", data={"connection_id": conn_id}, files={"file": ("jv.docx", p.read_bytes())}).json()
    assert r["tables"][0]["columns"] == ["Account", "Debit", "Credit"] and "Journal voucher" in r["text_preview"]
    xml = b"<rows><row><code>9001</code><name>Test Acc</name></row><row><code>9002</code><name>Test 2</name></row></rows>"
    r = client.post("/api/etl/extract", data={"connection_id": conn_id}, files={"file": ("acc.xml", xml)}).json()
    assert r["rows_extracted"] == 2


def test_load_records_api(client, conn_id):
    r = client.post("/api/etl/load-records", json={"connection_id": conn_id, "table": "accounts", "mode": "upsert",
                                                     "records": [{"code": "9999", "name": "API Test", "account_type": "expense"}]})
    assert r.status_code == 200, r.text
    assert r.json()["loaded"] == 1


ROWS = [("Invoice", "Customer", "Amount"), ("INV-1", "GKW Consult", "1200.00"),
        ("INV-2", "Acme Industries", "560.50"), ("INV-3", "Kite Foods", "99.00")]


def test_extract_image_pdf_pptx(tmp_path):
    import shutil
    from app.etl.extractors import extract
    if shutil.which("tesseract"):
        from PIL import Image, ImageDraw, ImageFont
        img = Image.new("RGB", (900, 260), "white")
        dr = ImageDraw.Draw(img)
        try:
            f = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 22)
        except OSError:
            f = None
        for i, r in enumerate(ROWS):
            for j, v in enumerate(r):
                dr.text((30 + j * 290, 20 + i * 55), v, fill="black", font=f)
        img.save(tmp_path / "scan.png")
        assert extract(tmp_path / "scan.png").tables[0]["df"].shape == (3, 3)
    try:
        from reportlab.platypus import SimpleDocTemplate, Table
        SimpleDocTemplate(str(tmp_path / "t.pdf")).build([Table([list(r) for r in ROWS])])
        assert extract(tmp_path / "t.pdf").tables[0]["df"].iloc[1]["Customer"] == "Acme Industries"
    except ImportError:
        pass
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    tbl = s.shapes.add_table(4, 3, Inches(1), Inches(1), Inches(6), Inches(2)).table
    for i, r in enumerate(ROWS):
        for j, v in enumerate(r):
            tbl.cell(i, j).text = v
    prs.save(tmp_path / "t.pptx")
    assert list(extract(tmp_path / "t.pptx").tables[0]["df"].columns) == ["Invoice", "Customer", "Amount"]
